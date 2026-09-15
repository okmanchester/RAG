import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from rag_baseline.chunk import run
from rag_baseline.parsers import parse_docx, parse_pdf
from rag_baseline.splitter import TokenCounter, chunk_document
from rag_baseline.structure import build_sections


def block(text, number=1, page=1, level=None, kind="paragraph"):
    return {"block_id": number, "text": text, "source": {"page": page},
            "heading_level": level, "kind": kind}


class ChunkingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.count = TokenCounter()

    def make(self, blocks, maximum=400):
        title, sections, _ = build_sections(blocks, "测试文档")
        doc = {"title": title, "doc_id": "test", "source_path": "test.pdf", "category": "test"}
        return chunk_document(doc, sections, self.count, maximum), sections

    def test_long_chinese_and_exact_coverage(self):
        value = ("第一句说明数据安全要求。第二句说明企业需要采取的措施；其中包含具体操作，和技术细节。\n\n" * 40).strip()
        chunks, _ = self.make([block(value)])
        self.assertGreater(len(chunks), 2)
        self.assertEqual("".join(c["body"][c["overlap_char_count"]:] for c in chunks), value)
        self.assertTrue(any(c["overlap_token_count"] > 0 for c in chunks[1:]))
        for c in chunks:
            self.assertEqual(self.count(c["text"]), c["token_count"])
            self.assertLessEqual(c["token_count"], 400)
            self.assertLessEqual(c["overlap_token_count"], 50)
            for s in c["sources"]:
                self.assertEqual(c["body"][s["body_char_start"]:s["body_char_end"]],
                                 value[s["block_char_start"]:s["block_char_end"]])

    def test_no_punctuation_and_unicode(self):
        value = "数据安全测试𠀀🙂" * 250
        chunks, _ = self.make([block(value)])
        self.assertEqual("".join(c["body"][c["overlap_char_count"]:] for c in chunks), value)
        self.assertTrue(all(c["body"].strip() for c in chunks))

    def test_nested_headings_cross_pages_and_no_cross_sections(self):
        blocks = [block("第一章 总则", 1), block("第一节 适用范围", 2), block("说明内容。", 3),
                  block("同章节跨页内容。", 4, 2), block("第二章 责任", 5, 2), block("其他规定。", 6, 2)]
        chunks, _ = self.make(blocks)
        self.assertEqual(chunks[0]["heading_path"], ["第一章 总则", "第一节 适用范围"])
        self.assertEqual({s["page"] for s in chunks[0]["sources"]}, {1, 2})
        self.assertEqual(chunks[1]["overlap_token_count"], 0)
        self.assertEqual(chunks[1]["heading_path"], ["第二章 责任"])

    def test_toc_does_not_leak_into_body(self):
        texts = ["中华人民共和国测试法", "目 录", "第一章 总则", "第二章 责任", "第一章 总则", "第一条 本法适用于测试。"]
        chunks, _ = self.make([block(t, i) for i, t in enumerate(texts, 1)])
        self.assertEqual(chunks[0]["heading_path"], ["目录"])
        self.assertIn("第二章 责任", chunks[0]["body"])
        self.assertEqual(chunks[1]["heading_path"], ["第一章 总则"])
        self.assertIn("第一条", chunks[1]["body"])

    def test_long_prose_is_not_heading(self):
        chunks, _ = self.make([block("（一）落实数据安全管理措施。企业应定期开展检查。", level=1)])
        self.assertEqual(chunks[0]["heading_path"], [])

    def test_title_rejects_order_number_navigation_and_generic_standard(self):
        for leading in ["第15号", "中华人民共和国国家标准", "首页 > 文件 > 通知"]:
            chunks, _ = self.make([block(leading, 1, level=1), block("测试文档", 2, level=2), block("正文。", 3)])
            self.assertEqual(chunks[0]["title"], "测试文档")

    def test_table_toc_and_page_edge_chapter(self):
        texts = [block("目 次", 1), block("|第一章 总则..........1|", 2, kind="table_row"),
                 block("|第二章 管理..........2|", 3, kind="table_row"),
                 block("第一章 总则", 4), block("第一条 正文。", 5)]
        chunks, _ = self.make(texts)
        self.assertEqual(chunks[-1]["heading_path"], ["第一章 总则"])
        value = "**第二章 信息服务规范**\n\n"
        with patch("pymupdf4llm.to_markdown", return_value=[{
            "metadata": {"page_number": 2}, "text": value + "2\n",
            "page_boxes": [{"class": "page-footer", "pos": [0, len(value)]},
                           {"class": "page-footer", "pos": [len(value), len(value) + 2]}],
        }]):
            blocks, metadata = parse_pdf(Path("test.pdf"))
            self.assertEqual(blocks[0]["text"], value.strip())
            self.assertEqual(len(blocks), 1)
            self.assertEqual(len(metadata["rescued_edge_headings"]), 1)

    def test_empty_no_heading_and_title_exhaustion(self):
        self.assertEqual(self.make([])[0], [])
        self.assertEqual(self.make([block("普通正文。")])[0][0]["heading_path"], [])
        with self.assertRaisesRegex(ValueError, "prefix"):
            self.make([block("普通正文。")], maximum=2)

    def test_deterministic_ids_and_document_isolation(self):
        chunks, sections = self.make([block("数据安全。" * 100)])
        self.assertEqual(chunks, self.make([block("数据安全。" * 100)])[0])
        doc = {"title": "测试文档", "doc_id": "other", "source_path": "other.pdf", "category": "test"}
        other = chunk_document(doc, sections, self.count)
        self.assertFalse({c["chunk_id"] for c in chunks} & {c["chunk_id"] for c in other})

    def test_docx_order_and_table(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "测试.docx"
            doc = Document()
            doc.add_paragraph("工业测试指南")
            doc.add_paragraph("一、总体要求")
            doc.add_paragraph("前置正文。")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "名称", "说明"
            table.cell(1, 0).text, table.cell(1, 1).text = "条目", "表格内容" * 200
            doc.add_paragraph("后置正文。")
            doc.save(path)
            blocks, _ = parse_docx(path)
            self.assertEqual([b["kind"] for b in blocks], ["paragraph"] * 3 + ["table_row"] * 2 + ["paragraph"])
            self.assertEqual(blocks[3]["source"]["row"], 1)
            self.assertNotIn("page", blocks[3]["source"])
            chunks, _ = self.make(blocks)
            self.assertTrue(all(c["heading_path"] == ["一、总体要求"] for c in chunks))
            self.assertTrue(all(c["token_count"] <= 400 for c in chunks))

    def test_pdf_parameters_and_empty_pages(self):
        with patch("pymupdf4llm.to_markdown", return_value=[
            {"metadata": {"page_number": 1}, "text": "# 标题\n\n正文"},
            {"metadata": {"page_number": 2}, "text": ""},
        ]) as mocked:
            blocks, metadata = parse_pdf(Path("test.pdf"))
            self.assertEqual(metadata["empty_pages"], [2])
            self.assertEqual(blocks[1]["source"]["page"], 1)
            self.assertFalse(mocked.call_args.kwargs["use_ocr"])
            self.assertFalse(mocked.call_args.kwargs["write_images"])

    def test_existing_text_in_formula_region_is_recovered(self):
        with patch("pymupdf4llm.to_markdown", return_value=[{
            "metadata": {"page_number": 1}, "text": "\n\n正文。",
            "page_boxes": [{"class": "formula", "pos": [0, 2], "bbox": [0, 0, 10, 10]}],
        }]), patch("rag_baseline.parsers.formula_text_regions", return_value={(1, 0): "第一章 总则"}):
            blocks, metadata = parse_pdf(Path("test.pdf"))
            self.assertEqual(blocks[0]["text"], "第一章 总则")
            self.assertEqual(len(metadata["recovered_formula_text"]), 1)
            self.assertEqual(self.make(blocks)[0][0]["heading_path"], ["第一章 总则"])

    def test_numeric_hierarchy_with_layout_spaces(self):
        chunks, _ = self.make([block("4 评价原则", 1), block("4 . 1 目标导向", 2), block("正文。", 3)])
        self.assertEqual(chunks[0]["heading_path"], ["4 评价原则", "4 . 1 目标导向"])

    def test_markdown_bullet_chapter_and_appendix_hierarchy(self):
        values = [block("9 日志要求", 1), block("旧章节。", 2), block("- **10** 安全防护要求", 3),
                  block("10.1 加密传输", 4), block("新章节。", 5), block("附录 B", 6, level=4),
                  block("（资料性）", 7, level=4), block("格式示例", 8, level=4),
                  block("B.1 属性格式", 9, level=5), block("B.1.1 静态格式", 10, level=5), block("表格文字。", 11)]
        chunks, _ = self.make(values)
        self.assertEqual(chunks[1]["heading_path"], ["10 安全防护要求", "10.1 加密传输"])
        self.assertEqual(chunks[-1]["heading_path"], ["附录 B", "格式示例", "B.1 属性格式", "B.1.1 静态格式"])

    def test_batch_continues_after_failure_and_reports_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, output = root / "raw", root / "out"
            source.mkdir()
            for name in ("a.pdf", "b.pdf", "c.pdf"):
                (source / name).write_bytes(b"fake")
            def fake_parse(path):
                if path.name == "a.pdf":
                    raise ValueError("broken PDF")
                return ([] if path.name == "b.pdf" else [block("有效文本。")], {"empty_pages": []})
            with patch("rag_baseline.chunk.parse", side_effect=fake_parse):
                self.assertEqual(run(source, output), 1)
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["summary"]["failed"], 1)
            self.assertEqual(report["summary"]["empty"], 1)
            self.assertEqual(report["summary"]["ok"], 1)


if __name__ == "__main__":
    unittest.main()
